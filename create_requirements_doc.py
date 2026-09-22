from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "AI悬疑推理游戏需求分析文档.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.35); sec.right_margin = Cm(2.35)

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.35
for name, size, space_before, space_after in [('Title', 24, 0, 12), ('Heading 1', 16, 18, 8), ('Heading 2', 12.5, 12, 5)]:
    s = styles[name]
    s.font.name = 'Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor(0,0,0)
    s.paragraph_format.space_before = Pt(space_before); s.paragraph_format.space_after = Pt(space_after)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def border_cell(cell, color='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = qn('w:' + edge); e = borders.find(tag)
        if e is None: e = OxmlElement('w:' + edge); borders.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),color)
def set_cell_text(cell, text, bold=False, color=None, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1); p.paragraph_format.line_spacing = 1.15
    r = p.add_run(str(text)); r.bold = bold; r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(8.5)
    if color: r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border_cell(cell)
def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,'1F4E78'); set_cell_text(c,h,True,(255,255,255),WD_ALIGN_PARAGRAPH.CENTER)
        if widths: c.width=Cm(widths[i])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            if ri%2==1: shade(cells[i],'F3F7FA')
            set_cell_text(cells[i],v,False,None, WD_ALIGN_PARAGRAPH.CENTER if i==0 else None)
            if widths: cells[i].width=Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def heading(text, level=1): doc.add_heading(text, level=level)
def para(text='', boldlead=None):
    p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Cm(0.74)
    if boldlead and text.startswith(boldlead):
        r=p.add_run(boldlead); r.bold=True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    return p
def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(text); return p
def page_break(): doc.add_page_break()

# Title page
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('AI 悬疑推理游戏需求分析文档')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(28); r=p.add_run('面向产品 研发 与测试评审'); r.font.size=Pt(13)
table(['文档信息','内容'],[
    ['项目名称','AI 悬疑推理游戏 Murder Mystery'],
    ['文档类型','需求分析'],
    ['版本','V1.0'],
    ['编写日期','2026 年 9 月 5 日'],
    ['适用对象','产品经理 前端开发 后端开发 算法工程师 测试人员'],
], [4,11])
para('本文以当前代码实现为基线，说明产品目标、范围、功能规则、数据与接口约束及验收口径。现有版本已具备完整的单人游戏闭环；文中标记为“后续建议”的内容不视为当前版本既有能力。')
page_break()

heading('1 项目概述')
para('AI 悬疑推理游戏是一款面向个人玩家的 Web 互动推理产品。玩家输入故事场景和人物数量后，系统调用大语言模型生成一宗独立的命案、角色关系和可供调查的线索。玩家以自然对话的方式与嫌疑人交流，整理线索，并在有限次数内指认真凶。')
heading('1.1 产品目标',2)
bullet('提供“每局内容不同”的轻量化剧本杀体验，降低传统固定剧本的重复游玩感。')
bullet('以自然闲聊替代问卷式审讯，让 AI 角色在不直接泄露真相的前提下给出符合人设的回答。')
bullet('保证从创建游戏、调查、指认到结算的完整且可理解的单人闭环。')
heading('1.2 范围与边界',2)
table(['范围内','当前不包含'],[
 ['单人网页端游戏','多人房间、实时匹配、语音或社交互动'],
 ['AI 生成案件、角色和对话','人工剧本编辑后台、剧本市场'],
 ['按会话保存运行中状态','跨设备云存档、账户体系和排行榜'],
 ['浏览器端本地历史记录','服务端持久化历史、数据分析后台'],
], [7.4,7.6])
heading('1.3 角色与核心诉求',2)
table(['角色','核心诉求','使用频率'],[
 ['玩家','快速开始一局可推理、可复盘的悬疑游戏','每局'],
 ['产品运营','控制体验边界，持续迭代内容质量与留存','版本迭代'],
 ['研发与测试','依据明确的状态、接口和验收规则实现与验证','开发测试期'],
], [3,9,3])

heading('2 用户旅程与业务流程')
para('玩家完成一局游戏的主流程如下。案件真相仅在后端状态中保留；游戏进行中，前端仅收到安全的角色与案情信息，结算时才揭示真凶和作案过程。')
table(['阶段','玩家行为','系统处理','产出'],[
 ['创建','填写场景与人物总数，开始游戏','生成角色、案件细节和开场旁白','案件概览、角色列表、初始线索'],
 ['调查','查看受害者资料或选择嫌疑人','生成卷宗简报或进入对应角色对话','调查简报、聊天记录、可核对线索'],
 ['辅助','点击 AI 帮我提问','结合当前角色、案情和最近对话生成一句自然追问，并取得回复','一轮辅助对话'],
 ['指认','选择嫌疑人并提交','校验编号并比对内部真凶；错误则扣减机会','胜利、继续调查或失败'],
 ['结算','查看结果后选择再来一局或返回首页','展示真凶与案情复盘，更新本地历史','游戏结果与历史记录'],
], [1.5,3.1,7.1,3.3])

heading('3 功能需求')
heading('3.1 游戏创建',2)
table(['编号','需求描述','验收标准'],[
 ['FR 01','首页提供故事场景文本输入与人物总数输入。','场景可为空并使用默认值；人物数可输入 3 至 8。'],
 ['FR 02','点击开始后建立 WebSocket 会话并初始化一局游戏。','前端进入游戏页，收到 game_init 后展示角色、案情和开场旁白。'],
 ['FR 03','系统生成一名受害者、一名真凶及若干嫌疑人。','人物数不超过玩家选择值；受害者和真凶始终保留。'],
 ['FR 04','生成与角色一致的死亡时间、地点、凶器、线索和作案过程。','案件概览字段非空；作案过程在未结算前不可对前端展示。'],
], [1.5,8.2,5.3])
heading('3.2 调查与角色对话',2)
table(['编号','需求描述','验收标准'],[
 ['FR 05','侧栏展示受害者与嫌疑人，受害者不得作为最终指认对象。','受害者显示为独立身份；指认列表仅含非受害者角色。'],
 ['FR 06','点击受害者后展示调查简报。','简报包括死因与时间、死前经历、旁人转述和 3 至 5 个疑点；不得点名或暗示真凶。'],
 ['FR 07','点击嫌疑人后进入该角色对话，显示人物背景和既有消息。','系统提示以自然聊天方式交流；切换或退出后恢复调查状态。'],
 ['FR 08','玩家可发送非空文本；系统以当前角色人设和案件信息生成回复。','玩家消息与 AI 回复按顺序显示；回复不主动承认真凶身份或交代决定性破绽。'],
 ['FR 09','玩家可使用 AI 帮我提问。','系统基于最近对话生成一句自然问题，并自动完成该轮问答。'],
 ['FR 10','玩家可退出对话。','发送 EXIT 后清空当前选择，提示可继续调查或指认。'],
], [1.5,8.2,5.3])
heading('3.3 线索与指认',2)
table(['编号','需求描述','验收标准'],[
 ['FR 11','案件初始线索在侧栏按常见分隔符拆分展示，支持本局置顶。','玩家点击线索可切换置顶状态；置顶仅影响当前页面展示。'],
 ['FR 12','玩家可随时打开最终指认面板。','进入指认时退出当前角色对话；支持点击或输入嫌疑人编号。'],
 ['FR 13','每局初始有 3 次有效指认机会。','正确指认立即胜利；错误指认扣 1 次；无效编号不扣次数；耗尽后失败。'],
 ['FR 14','结算页展示胜负、真凶和案件复盘。','无论胜负均展示 killer_name；作案过程仅在结算页可见。'],
], [1.5,8.2,5.3])
heading('3.4 历史与再次游戏',2)
table(['编号','需求描述','验收标准'],[
 ['FR 15','浏览器记录已开始局的轮次、场景、时间和胜负。','刷新后本地记录仍可读取；未完成上一局开启新局时，上一局标为失败。'],
 ['FR 16','结算后支持用相同参数再来一局，或返回首页重新配置。','再来一局创建新轮次；返回首页断开 WebSocket 并重置当前状态。'],
], [1.5,8.2,5.3])

page_break()
heading('4 关键业务规则')
table(['规则编号','规则','说明'],[
 ['BR 01','真相隔离','内部角色可标记为 Killer；对外序列化后除 Victim 外均显示为 Suspect，且游戏中移除 murder_process。'],
 ['BR 02','受害者调查','只有点击受害者时才生成受害者调查简报；简报生成失败时给出可读的兜底内容。'],
 ['BR 03','对话风格','NPC 应像熟人闲聊，可回避、反问、隐瞒或说谎，但回答需与人设和案件保持基本自洽。'],
 ['BR 04','指认编号','前端编号以嫌疑人列表为准，从 0 映射到内部角色数组；不得把受害者计入编号。'],
 ['BR 05','会话隔离','不同 session_id 的运行中状态彼此隔离；断开连接后未启动会话可清理。'],
 ['BR 06','结果不可逆','正确指认或次数耗尽后进入游戏结束状态，前端不应再允许继续调查或再次提交。'],
], [2.1,4.0,8.9])

heading('5 数据与接口需求')
heading('5.1 核心数据对象',2)
table(['对象','关键字段','用途'],[
 ['Character','role, name, backstory, relation_to_victim','角色身份、背景与死者关系；内部 role 可为 killer。'],
 ['StoryDetails','victim_name, time_of_death, location_found, murder_weapon, initial_clues, murder_process','案件概览、线索与仅结算可见的真相。'],
 ['GameState','environment, max_characters, characters, messages, selected_character_id, num_guesses_left, result','单局运行时状态。'],
 ['GameHistoryRecord','round, result, timestamp, environment, killer_name','浏览器端本地历史记录。'],
], [3.1,7.2,4.7])
heading('5.2 通信接口',2)
para('后端提供 GET /health 健康检查，并通过 /ws/game/{session_id} 承担一局游戏的双向通信。消息采用 JSON。')
table(['方向','消息或动作','关键参数','预期结果'],[
 ['客户端→服务端','start_game','environment, max_characters','创建案件并返回 game_init'],
 ['客户端→服务端','inspect_victim','char_id','返回含 victim_report 的状态'],
 ['客户端→服务端','select_character','char_id 或 null','选择角色或进入指认模式'],
 ['客户端→服务端','player_message','text','记录消息并返回 NPC 回复；EXIT 用于退出'],
 ['客户端→服务端','ask_sherlock','无','返回 AI 建议问题及 NPC 回复'],
 ['客户端→服务端','make_guess','guess_idx','返回 state_update 或 game_over'],
 ['服务端→客户端','game_init / state_update / game_over','data: GameState','下发安全状态；game_over 额外含 killer_name'],
], [2.6,3.6,4.8,4.0])

page_break()
heading('6 非功能需求')
table(['类别','需求','验收口径'],[
 ['安全与公平','游戏进行中不得泄露真凶身份或完整作案过程。','检查所有非 game_over 下发数据：角色角色字段不含 Killer，story_details 不含 murder_process。'],
 ['稳定性','大模型不可用、超时或返回非预期结构时，应有明确失败提示或可控兜底。','API Key 未配置时开始游戏返回可定位错误；结构化生成具备校验与重试策略。'],
 ['性能','界面操作应即时反馈；长耗时 AI 调用不应造成浏览器无响应。','建议补充加载状态与超时提示；后端连接、读取超时可配置。'],
 ['可用性','游戏核心操作应无需阅读技术说明即可理解。','玩家可在首页开始、在游戏页调查和指认、在结算页复盘与重开。'],
 ['兼容性','支持现代桌面浏览器。','验证 Chrome、Edge 的 WebSocket 与 localStorage 行为。'],
 ['可维护性','前后端状态模型与消息类型保持一致。','类型定义、Pydantic 模型和接口动作变更时同步更新测试。'],
], [2.4,7.5,5.1])

heading('7 异常与边界场景')
table(['场景','期望处理'],[
 ['未开始游戏即发送调查、对话或指认动作','服务端忽略请求，不产生异常状态。'],
 ['受害者或嫌疑人编号无效','调查请求忽略；指认请求提示有效编号范围且不扣次数。'],
 ['未选择角色即发送消息','保留玩家消息并提示先选择角色。'],
 ['空白输入','前端不发送，后端也不执行对话生成。'],
 ['WebSocket 未连接或仍在连接','客户端缓存动作，连接成功后依次发送。'],
 ['玩家在进行中直接开始新局','历史中上一局标记为失败，新局使用新的轮次。'],
 ['大模型响应异常','建议统一返回可读错误、允许重新尝试，并保留已生成的安全状态。'],
], [5.2,9.8])

heading('8 测试验收清单')
table(['编号','测试点','通过标准'],[
 ['AC 01','创建默认场景的一局游戏','展示案情、至少一名受害者和若干嫌疑人，且无真凶泄露。'],
 ['AC 02','创建不同场景与人物数','人物数符合范围和配置，案件字段完整可展示。'],
 ['AC 03','查看受害者资料','显示调查简报且未出现任何具体嫌疑人指向。'],
 ['AC 04','与嫌疑人对话并使用 AI 提问','聊天记录正确追加；提问与回答围绕当前角色和案情。'],
 ['AC 05','退出对话后指认','退出后显示指认面板；受害者不在候选列表。'],
 ['AC 06','提交无效、错误和正确指认','无效不扣次数；错误扣一次并可继续；正确立即进入胜利结算。'],
 ['AC 07','连续三次错误指认','第三次后失败结算，并展示真凶与复盘。'],
 ['AC 08','再来一局与返回首页','新局轮次递增；返回首页后当前状态被重置。'],
 ['AC 09','刷新并查看历史','历史记录保留轮次、时间、场景和结果。'],
], [1.5,7.0,6.5])

heading('9 后续迭代建议')
para('以下内容用于规划后续版本，不属于当前验收范围。优先级应以内容质量和核心闭环稳定性为先。')
table(['优先级','建议','预期价值'],[
 ['P0','增加生成中、对话中和失败重试的状态提示；统一大模型异常反馈。','降低等待焦虑，避免玩家误判页面失效。'],
 ['P0','为案件生成增加一致性校验，例如人物、时间线、凶器与线索可交叉验证。','减少无法推理或自相矛盾的案件。'],
 ['P1','支持线索笔记、标签与跨角色对话摘要。','帮助玩家组织推理过程，提升中长局体验。'],
 ['P1','增加服务端会话存档和恢复机制。','支持刷新或断线后继续当前局。'],
 ['P2','提供难度、题材、时长等配置，以及人工精选剧本。','扩大可玩性并支持内容运营。'],
 ['P2','引入账户、成就、排行榜或多人协作。','提升长期留存和社交性。'],
], [1.7,8.3,5.0])

heading('10 结论')
para('当前项目已经覆盖“创建案件 - 角色调查 - AI 对话 - 有限指认 - 结果复盘 - 本地历史”的最小可用产品闭环。下一阶段应优先完善生成结果的一致性、调用过程的可见性和异常恢复，再扩展存档、难度与内容运营能力。')

# Footer page number
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('AI 悬疑推理游戏需求分析文档  |  '); r.font.size=Pt(8); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)

doc.core_properties.title='AI 悬疑推理游戏需求分析文档'
doc.core_properties.subject='需求分析'
doc.core_properties.author='Codex'
doc.save(OUT)
print(OUT)
